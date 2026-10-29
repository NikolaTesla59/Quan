# -*- coding: utf-8 -*-
"""
Script to build the official Executive Presentation Paper:
Bao_cao_Thuyet_trinh_Agribank_MB_2020-2025_Executive.docx
"""
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from copy import deepcopy
import os

source_v5_path = r'C:\Users\Admin\Desktop\NEU (ThS)\Quản trị rủi ro\files 2\Tieu_luan_Agribank_MB_2020-2025_hoan_thien_v5.docx'
target_exec_path = r'C:\Users\Admin\Desktop\NEU (ThS)\Quản trị rủi ro\files 2\Bao_cao_Thuyet_trinh_Agribank_MB_2020-2025_Executive.docx'
charts_dir = r'C:\Users\Admin\.gemini\antigravity\scratch\charts'

doc_src = docx.Document(source_v5_path)
body_src = doc_src._body._element

doc_dst = docx.Document()

# 1. Page Setup: Standard 1.0 inch margins, Letter size
for section in doc_dst.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)

# 2. Settings: Enable automatic dynamic field updates
settings = doc_dst.settings.element
update_fields = parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>')
settings.append(update_fields)

# Helper formatting functions
def style_run(run, font_name="Times New Roman", size_pt=11, bold=False, italic=False, color_rgb=(0,0,0)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

def add_p(doc, text="", space_before=2, space_after=3, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT, page_break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.page_break_before = page_break_before
    p.alignment = align
    if text:
        r = p.add_run(text)
        style_run(r, font_name="Times New Roman", size_pt=11)
    return p

def add_h1(doc, text, page_break_before=True):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = page_break_before
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    style_run(r, font_name="Times New Roman", size_pt=13.5, bold=True, color_rgb=(0, 32, 96))
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    style_run(r, font_name="Times New Roman", size_pt=12, bold=True, color_rgb=(15, 76, 129))
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    style_run(r, font_name="Times New Roman", size_pt=11, bold=True, italic=True, color_rgb=(0, 0, 0))
    return p

def add_dot_leader_tab(paragraph, pos=9360):
    pPr = paragraph._element.get_or_add_pPr()
    tabs = parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:leader="dot" w:pos="{pos}"/></w:tabs>')
    pPr.append(tabs)

def add_image_exhibit(doc, chart_filename, title_text, source_text):
    chart_path = os.path.join(charts_dir, chart_filename)
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.keep_with_next = True
    r_t = p_title.add_run(title_text)
    style_run(r_t, font_name="Times New Roman", size_pt=10.5, bold=True, color_rgb=(0, 32, 96))
    
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(4)
    p_img.paragraph_format.space_after = Pt(4)
    r_img = p_img.add_run()
    r_img.add_picture(chart_path, width=Inches(6.2))
    
    p_src = doc.add_paragraph()
    p_src.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_src.paragraph_format.space_before = Pt(2)
    p_src.paragraph_format.space_after = Pt(10)
    r_s = p_src.add_run(source_text)
    style_run(r_s, font_name="Times New Roman", size_pt=9.0, italic=True, color_rgb=(80, 80, 80))

print("Base setup configured successfully.")
