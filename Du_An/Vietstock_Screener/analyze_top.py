import openpyxl
import pandas as pd

wb = openpyxl.load_workbook('Bo_Loc_Co_Phieu_Vietstock_Khung_Tieu_Chi_PRO.xlsx', data_only=True)
s = wb['Top_Co_Phieu_Chon_Loc']
headers = [s.cell(5, c).value for c in range(1, s.max_column+1)]

data = []
for r in range(6, s.max_row+1):
    vals = [s.cell(r, c).value for c in range(1, s.max_column+1)]
    if vals[1]:
        data.append(vals)

df_top = pd.DataFrame(data, columns=headers)

print("="*110)
print(f"{'Hạng':<5}{'Mã':<6}{'Tên Doanh Nghiệp':<30}{'Giá (đ)':<10}{'Target (đ)':<12}{'Upside':<10}{'P/E':<8}{'Fwd P/E':<10}{'ROE':<8}{'Khuyến Nghị':<30}")
print("="*110)

for _, r in df_top.head(25).iterrows():
    h = int(r['Hạng']) if r['Hạng'] else 0
    ma = str(r['Mã CK'])
    ten = str(r['Tên Doanh Nghiệp'])[:28]
    gia = f"{float(r['Thị Giá (VNĐ)']):,.0f}" if r['Thị Giá (VNĐ)'] else "-"
    target = f"{float(r['Giá Mục Tiêu (VNĐ)']):,.0f}" if r['Giá Mục Tiêu (VNĐ)'] else "-"
    upside_val = float(r['Upside Kỳ Vọng (%)']) * 100 if r['Upside Kỳ Vọng (%)'] else 0
    upside = f"{upside_val:+.1f}%"
    pe = f"{float(r['P/E Hiện Tại']):.2f}" if r['P/E Hiện Tại'] else "-"
    fpe = f"{float(r['Forward P/E']):.2f}" if r['Forward P/E'] else "-"
    roe_val = float(r['ROE (%)']) * 100 if r['ROE (%)'] and float(r['ROE (%)']) < 1 else (float(r['ROE (%)']) if r['ROE (%)'] else 0)
    roe = f"{roe_val:.1f}%"
    kn = str(r['Khuyến Nghị Định Chế'])
    print(f"{h:<5}{ma:<6}{ten:<30}{gia:<10}{target:<12}{upside:<10}{pe:<8}{fpe:<10}{roe:<8}{kn:<30}")

print("\n" + "="*110)
print("CHI TIẾT LUẬN ĐIỂM CÁC MÃ DẪN ĐẦU THEO CÁC NHÓM TIÊU CHÍ:")
print("="*110)
for _, r in df_top.head(10).iterrows():
    print(f"[{r['Mã CK']}] - {r['Tên Doanh Nghiệp']} ({r['Ngành']})")
    print(f"  • Chiến lược: {r['Chiến Lược']} | Tổng điểm: {r['Tổng Điểm']}/100")
    print(f"  • Định giá: Thị giá {float(r['Thị Giá (VNĐ)']):,.0f} đ -> Target {float(r['Giá Mục Tiêu (VNĐ)']):,.0f} đ (Upside: {float(r['Upside Kỳ Vọng (%)'])*100:.1f}%)")
    print(f"  • Chỉ số cốt lõi: P/E: {r['P/E Hiện Tại']} | Fwd P/E: {r['Forward P/E']} | ROE: {float(r['ROE (%)'])*100:.1f}% | EV/EBITDA: {r['EV/EBITDA']}")
    print(f"  • Luận điểm & Xúc tác: {r['Luận Điểm Trọng Tâm Theo Khung & Xúc Tác']}")
    print("-" * 100)
