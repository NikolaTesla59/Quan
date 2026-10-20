import openpyxl
import pandas as pd

wb = openpyxl.load_workbook('Bo_Loc_Co_Phieu_Vietstock_Khung_Tieu_Chi_PRO.xlsx', data_only=True)
s = wb['Top_Co_Phieu_Chon_Loc']
headers = [s.cell(5, c).value for c in range(1, s.max_column+1)]
data = []
for r in range(6, s.max_row+1):
    vals = [s.cell(r, c).value for c in range(1, s.max_column+1)]
    if vals[1]:
        data.append(vals)
df = pd.DataFrame(data, columns=headers)

df['PE'] = pd.to_numeric(df['P/E Hiện Tại'], errors='coerce')
df['FwdPE'] = pd.to_numeric(df['Forward P/E'], errors='coerce')
df['ROE'] = pd.to_numeric(df['ROE (%)'], errors='coerce')
df['ROE_pct'] = df['ROE'].apply(lambda x: x*100 if x < 1 else x)
df['Upside'] = pd.to_numeric(df['Upside Kỳ Vọng (%)'], errors='coerce') * 100
df['Gia'] = pd.to_numeric(df['Thị Giá (VNĐ)'], errors='coerce')
df['Target'] = pd.to_numeric(df['Giá Mục Tiêu (VNĐ)'], errors='coerce')
df['Diem'] = pd.to_numeric(df['Tổng Điểm'], errors='coerce')

print("=== 1. TOP UPSIDE ĐỊNH GIÁ (Điểm >= 78) ===")
sub_up = df[df['Diem'] >= 78].sort_values(by='Upside', ascending=False).head(10)
for _, r in sub_up.iterrows():
    print(f"{r['Mã CK']:<6} | {r['Tên Doanh Nghiệp'][:28]:<28} | Thị giá: {r['Gia']:>8,.0f} đ | Target: {r['Target']:>8,.0f} đ | Upside: {r['Upside']:>6.1f}% | P/E: {r['PE']:>5.2f} | ROE: {r['ROE_pct']:>5.1f}% | Điểm: {r['Diem']}")

print("\n=== 2. TOP P/E THẤP + ROE CAO (ROE >= 18%, P/E <= 8, Điểm >= 78) ===")
sub_pe = df[(df['Diem'] >= 78) & (df['ROE_pct'] >= 18) & (df['PE'] > 0) & (df['PE'] <= 8)].sort_values(by='PE', ascending=True).head(10)
for _, r in sub_pe.iterrows():
    print(f"{r['Mã CK']:<6} | {r['Tên Doanh Nghiệp'][:28]:<28} | P/E: {r['PE']:>5.2f} | Fwd P/E: {r['FwdPE']:>5.2f} | ROE: {r['ROE_pct']:>5.1f}% | Upside: {r['Upside']:>6.1f}% | Thị giá: {r['Gia']:>8,.0f} đ")

print("\n=== 3. TOP ROE DẪN ĐẦU (ROE cao nhất, Điểm >= 78) ===")
sub_roe = df[df['Diem'] >= 78].sort_values(by='ROE_pct', ascending=False).head(10)
for _, r in sub_roe.iterrows():
    print(f"{r['Mã CK']:<6} | {r['Tên Doanh Nghiệp'][:28]:<28} | ROE: {r['ROE_pct']:>5.1f}% | P/E: {r['PE']:>5.2f} | Fwd P/E: {r['FwdPE']:>5.2f} | Upside: {r['Upside']:>6.1f}% | Thị giá: {r['Gia']:>8,.0f} đ")
